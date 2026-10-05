import streamlit as st
import pandas as pd
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import unicodedata
import re

# =========================================================================
# CONFIGURAÇÃO DA PÁGINA
# =========================================================================
st.set_page_config(page_title="Cherry Picking Universal - Natura", page_icon="💵", layout="wide")

st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Poppins:wght@300;400;500;600;700&display=swap');
    html, body, [data-testid="stAppViewContainer"], .stApp { font-family: 'Poppins', sans-serif !important; }
    </style>
    """, unsafe_allow_html=True
)

st.title("💵 Agente de Sourcing: Cherry Picking Universal")
st.write("Suba planilhas de **qualquer categoria** (Mão de Obra, Frota, Serviços, Peças) para consolidação automática.")

# Inicialização da memória de sessão do Streamlit
if "df_consolidado_final" not in st.session_state:
    st.session_state.df_consolidado_final = None
if "buffer_excel_final" not in st.session_state:
    st.session_state.buffer_excel_final = None
if "fornecedores_processados" not in st.session_state:
    st.session_state.fornecedores_processados = []

arquivos_carregados = st.file_uploader(
    "Suba aqui as cotações (.xlsx ou .csv):", 
    type=["xlsx", "csv"], 
    accept_multiple_files=True
)

# =========================================================================
# FUNÇÕES DE HIGIENIZAÇÃO DE DADOS
# =========================================================================
def padronizar_texto(texto):
    if pd.isna(texto) or texto is None:
        return ""
    texto = str(texto).upper().strip()
    texto = ''.join(c for c in unicodedata.normalize('NFD', texto) if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', texto)

def limpar_valor(val):
    if pd.isna(val) or val is None:
        return None
    try:
        if isinstance(val, (int, float)):
            return float(val) if val > 0 else None
        
        texto = str(val).strip().upper()
        numeros = re.findall(r'[0-9\.,\-]+', texto)
        if not numeros: return None
        
        texto_limpo = numeros[0]
        if "." in texto_limpo and "," in texto_limpo:
            texto_limpo = texto_limpo.replace(".", "").replace(",", ".")
        elif "," in texto_limpo:
            texto_limpo = texto_limpo.replace(",", ".")
            
        val_float = float(texto_limpo)
        return val_float if val_float > 0 else None
    except Exception:
        return None

# =========================================================================
# LEITURA INTELIGENTE E ADAPTÁVEL
# =========================================================================
def processar_planilha_generica(file_obj, nome_arquivo):
    dfs_extraidos = []
    
    if nome_arquivo.endswith(".csv"):
        df_raw = pd.read_csv(file_obj)
        dfs_extraidos.append((nome_arquivo, df_raw))
    else:
        xl = pd.ExcelFile(file_obj)
        for sheet in xl.sheet_names:
            try:
                df_raw = pd.read_excel(file_obj, sheet_name=sheet, header=None)
            except Exception:
                continue
                
            if df_raw.empty or len(df_raw) < 2: 
                continue

            # Localiza a linha de cabeçalho buscando por qualquer palavra descritiva ou financeira
            header_row = 0
            max_matches = 0
            
            kw_gerais = [
                "regi", "cd", "local", "rota", "origem", "destino", "uf", "cidade", "unidade", "site",
                "carg", "funç", "func", "veic", "veículo", "item", "descri", "perfil", "linha",
                "turn", "horar", "jornada",
                "valor", "preço", "preco", "custo", "diaria", "diária", "total", "tarifa", "frete", "r$", "faturamento"
            ]
            
            for r_idx in range(min(30, len(df_raw))):
                row_vals = [padronizar_texto(v) for v in df_raw.iloc[r_idx].values if pd.notna(v)]
                matches = sum(1 for val in row_vals for kw in kw_gerais if kw.upper() in val)
                if matches > max_matches and matches >= 2:
                    max_matches = matches
                    header_row = r_idx

            df_sheet = pd.read_excel(file_obj, sheet_name=sheet, header=header_row)
            label = f"{nome_arquivo} ({sheet.strip()})" if len(xl.sheet_names) > 1 else nome_arquivo
            dfs_extraidos.append((label, df_sheet))

    dfs_normalizados = []

    for label, df in dfs_extraidos:
        df = df.loc[:, ~df.columns.astype(str).str.contains('^Unnamed')].dropna(how='all')
        if df.empty: continue
            
        col_desc1, col_desc2, col_valor = None, None, None
        
        # Identificação flexível da 1ª e 2ª coluna descritiva
        for col in df.columns:
            c_lower = str(col).strip().lower()
            if any(k in c_lower for k in ["regi", "cd", "local", "rota", "origem", "destino", "uf", "unidade", "item", "carg", "veic"]):
                if not col_desc1:
                    col_desc1 = col
                elif not col_desc2 and col != col_desc1:
                    col_desc2 = col

        # Identificação da coluna de valor
        for col in df.columns:
            c_lower = str(col).strip().lower()
            if any(k in c_lower for k in ["faturamento", "valor total", "custo total", "preço final", "preco final", "total"]):
                col_valor = col
                break
                
        if not col_valor:
            for col in df.columns:
                c_lower = str(col).strip().lower()
                if any(k in c_lower for k in ["valor", "preço", "preco", "diaria", "tarifa", "r$"]) and not any(bad in c_lower for bad in ["%", "taxa", "imposto"]):
                    col_valor = col
                    break

        # Fallback de preço: primeira coluna com números
        if not col_valor:
            for col in df.columns:
                if col in [col_desc1, col_desc2]: continue
                vals_col = df[col].apply(limpar_valor).dropna()
                if len(vals_col) > 0 and vals_col.mean() > 0.5:
                    col_valor = col
                    break

        if not col_desc1 and len(df.columns) > 0:
            col_desc1 = df.columns[0]

        if col_desc1 and col_valor:
            df_sub = pd.DataFrame()
            df_sub["Local / Origem / Região"] = df[col_desc1].apply(padronizar_texto)
            df_sub["Item / Rota / Perfil"] = df[col_desc2].apply(padronizar_texto) if col_desc2 else "GERAL"
            df_sub[label] = df[col_valor].apply(limpar_valor)
            
            df_sub = df_sub.dropna(subset=["Local / Origem / Região", label])
            df_sub = df_sub[~df_sub["Local / Origem / Região"].isin(["", "NAN", "TOTAL", "SUBTOTAL"])]
            
            if not df_sub.empty:
                dfs_normalizados.append(df_sub)

    return dfs_normalizados

# =========================================================================
# ESTILIZAÇÃO DO EXCEL
# =========================================================================
def estilizar_planilha_excel(df, fornecedores):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cherry Picking"
    
    colunas_chave = ["Local / Origem / Região", "Item / Rota / Perfil"]
    colunas_completas = colunas_chave + fornecedores + ["Melhor Preço", "Fornecedor Vencedor"]
    ws.append(colunas_completas)
    
    idx_min_col = len(colunas_chave) + len(fornecedores) + 1
    idx_winner_col = idx_min_col + 1
    total_cols = idx_winner_col
    
    header_fill = PatternFill(start_color="9E472A", end_color="9E472A", fill_type="solid")
    header_font = Font(name="Poppins", size=10, bold=True, color="FFFFFF")
    border_cinza = Border(left=Side(style='thin', color='E0D8D3'), right=Side(style='thin', color='E0D8D3'), top=Side(style='thin', color='E0D8D3'), bottom=Side(style='thin', color='E0D8D3'))
    
    for col_idx in range(1, total_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_cinza

    row_num = 2
    for _, row in df.iterrows():
        ws.cell(row=row_num, column=1, value=str(row.get("Local / Origem / Região", "")).title())
        ws.cell(row=row_num, column=2, value=str(row.get("Item / Rota / Perfil", "")).title())
        
        for col_idx, forn in enumerate(fornecedores, start=3):
            ws.cell(row=row_num, column=col_idx, value=row.get(forn, None))
        
        col_let_start = get_column_letter(3)
        col_let_end = get_column_letter(2 + len(fornecedores))
        col_let_min = get_column_letter(idx_min_col)
        
        ws.cell(row=row_num, column=idx_min_col, value=f"=MIN({col_let_start}{row_num}:{col_let_end}{row_num})")
        ws.cell(row=row_num, column=idx_winner_col, value=f'=INDEX(${col_let_start}$1:${col_let_end}$1, MATCH({col_let_min}{row_num}, {col_let_start}{row_num}:{col_let_end}{row_num}, 0))')
        
        coluna_vencedora_python = row.get("Fornecedor Vencedor")
        col_idx_vencedora = fornecedores.index(coluna_vencedora_python) + 3 if coluna_vencedora_python in fornecedores else None
        
        winner_fill = PatternFill(start_color="E6F0EA", end_color="E6F0EA", fill_type="solid")
        winner_font = Font(name="Poppins", size=10, bold=True, color="1E3D2F")
        normal_font = Font(name="Poppins", size=10, color="000000")
        
        for col_idx in range(1, total_cols + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.border = border_cinza
            cell.font = normal_font
            
            if col_idx == col_idx_vencedora:
                cell.fill = winner_fill
                cell.font = winner_font
                
            if col_idx >= 3:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if col_idx <= idx_min_col:
                    cell.number_format = 'R$ #,##0.00'
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        row_num += 1

    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = 25
        
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

# =========================================================================
# INTERFACE E FLUXO PRINCIPAL
# =========================================================================
if arquivos_carregados:
    st.info(f"📁 **{len(arquivos_carregados)} arquivo(s)** prontos para processamento.")
    
    if st.button("🚀 Processar Cherry Picking"):
        dfs_processados = []
        relatorio_erros = []
        
        progress_bar = st.progress(0)
        status_text = st.empty()
        
        for idx, arquivo in enumerate(arquivos_carregados):
            nome_clean = arquivo.name.split(".")[0].replace("Template_para_Precificação_-_", "").replace("Cópia_de_", "").strip()
            status_text.text(f"⏳ Lendo arquivo {idx+1} de {len(arquivos_carregados)}: {arquivo.name}...")
            
            try:
                sub_dfs = processar_planilha_generica(arquivo, nome_clean)
                if sub_dfs:
                    dfs_processados.extend(sub_dfs)
                else:
                    relatorio_erros.append((arquivo.name, "Sem colunas válidas de descrição/valor."))
            except Exception as e:
                relatorio_erros.append((arquivo.name, f"Erro: {str(e)}"))
                
            progress_bar.progress((idx + 1) / len(arquivos_carregados))
            
        status_text.empty()
        progress_bar.empty()
        
        if relatorio_erros:
            with st.expander("⚠️ Arquivos com divergência ou ignorados"):
                for arq, msg in relatorio_erros:
                    st.warning(f"**{arq}**: {msg}")

        if dfs_processados:
            colunas_fornecedores = [list(d.columns)[-1] for d in dfs_processados]
            
            df_consolidado = dfs_processados[0]
            for df_prox in dfs_processados[1:]:
                df_consolidado = pd.merge(
                    df_consolidado, df_prox, 
                    on=["Local / Origem / Região", "Item / Rota / Perfil"], 
                    how="outer"
                )

            df_consolidado = df_consolidado.sort_values(by=["Local / Origem / Região"]).reset_index(drop=True)

            def calc_melhor_preco(row):
                vals = [row[c] for c in colunas_fornecedores if pd.notna(row[c]) and row[c] is not None]
                return min(vals) if vals else None

            def calc_fornecedor_vencedor(row):
                best = calc_melhor_preco(row)
                if best is None: return None
                for c in colunas_fornecedores:
                    if row[c] == best: return c
                return None

            df_consolidado["Melhor Preço"] = df_consolidado.apply(calc_melhor_preco, axis=1)
            df_consolidado["Fornecedor Vencedor"] = df_consolidado.apply(calc_fornecedor_vencedor, axis=1)

            # SALVA NA SESSÃO DO STREAMLIT PARA NÃO SUMIR AO REFRESH
            st.session_state.df_consolidado_final = df_consolidado
            st.session_state.buffer_excel_final = estilizar_planilha_excel(df_consolidado, colunas_fornecedores)
            st.session_state.fornecedores_processados = colunas_fornecedores
            
            st.balloons()
            st.success("✨ Consolidação concluída com sucesso!")

# SE OS DADOS JÁ FORAM PROCESSADOS, EXIBE NA TELA SEM SUMIR
if st.session_state.df_consolidado_final is not None:
    st.write("📊 **Prévia da Tabela Consolidada:**")
    st.dataframe(st.session_state.df_consolidado_final, use_container_width=True)

    st.download_button(
        label="📥 Baixar Consolidação em Excel (.xlsx)",
        data=st.session_state.buffer_excel_final,
        file_name="Cherry_Picking_Consolidado_Final.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
